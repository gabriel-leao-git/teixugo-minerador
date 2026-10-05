"""Renderizadores de Word, Excel e PDF. Dependências OPCIONAIS (importadas só quando usadas):

    python-docx  -> .docx        openpyxl  -> .xlsx        reportlab  -> .pdf

Instale com:  pip install -r requirements-optional.txt
"""
from __future__ import annotations

from typing import Any

from . import economics
from .blocks import collect_sources, evidence_rows, strings
from .model import FIELD_KEYS, fmt_value, scope_of

TAG_RGB = {"verified": (0x1A, 0x7F, 0x37), "estimated": (0x9A, 0x67, 0x00), "unverified": (0x6E, 0x77, 0x81)}


class MissingDependency(RuntimeError):
    """A biblioteca opcional necessária para o formato não está instalada."""


def _need(module: str, package: str, fmt: str):
    try:
        return __import__(module, fromlist=["_"])
    except ImportError as exc:
        raise MissingDependency(
            f"Para gerar .{fmt} instale '{package}':  pip install {package}  (ou pip install -r requirements-optional.txt)"
        ) from exc


# --------------------------------------------------------------------------- DOCX


def to_docx(blocks: list[dict[str, Any]], path: str) -> None:
    _need("docx", "python-docx", "docx")
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    for b in blocks:
        t = b["t"]
        if t == "title":
            doc.add_heading(b["text"], 0)
        elif t == "banner":
            r = doc.add_paragraph().add_run("⚠ " + b["text"])
            r.bold = True
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x00)
        elif t == "meta":
            doc.add_paragraph("   ·   ".join(f"{k}: {v}" for k, v in b["items"]))
        elif t == "h2":
            doc.add_heading(b["text"], 1)
        elif t == "h3":
            doc.add_heading(b["text"], 2)
        elif t == "p":
            doc.add_paragraph(b["text"])
        elif t == "bullets":
            for x in b["items"]:
                doc.add_paragraph(str(x), style="List Bullet")
        elif t == "ol":
            for x in b["items"]:
                doc.add_paragraph(str(x), style="List Number")
        elif t == "note":
            r = doc.add_paragraph().add_run(b["text"])
            r.italic = True
            r.font.size = Pt(8.5)
        elif t == "field":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"{b['label']}: ").bold = True
            p.add_run(str(b["value"]))
            if b.get("tag"):
                tr = p.add_run(f"  [{b['tagtext']}]")
                tr.bold = True
                tr.font.size = Pt(8.5)
                tr.font.color.rgb = RGBColor(*TAG_RGB[b["tag"]])
            if b.get("src"):
                sr = p.add_run(f"  ({b['src']})")
                sr.font.size = Pt(8)
                sr.font.color.rgb = RGBColor(0x6E, 0x77, 0x81)
        elif t == "table":
            table = doc.add_table(rows=1, cols=len(b["headers"]))
            table.style = "Table Grid"
            for i, h in enumerate(b["headers"]):
                cell = table.rows[0].cells[i]
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(h))
                run.bold = True
                run.font.size = Pt(8.5)
            for row in b["rows"]:
                cells = table.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = ""
                    cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(8.5)
            doc.add_paragraph()
    doc.save(path)


# --------------------------------------------------------------------------- XLSX


def to_xlsx(report: dict[str, Any], path: str) -> None:
    _need("openpyxl", "openpyxl", "xlsx")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    meta = report["meta"]
    S = strings(meta["language"])
    products = report["products"]
    wb = Workbook()

    head_fill = PatternFill("solid", fgColor="1F2937")
    head_font = Font(bold=True, color="FFFFFF")

    def sheet(title: str, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for i in range(len(headers)):
            ws.column_dimensions[get_column_letter(i + 1)].width = (widths[i] if widths else 22)
        ws.freeze_panes = "A2"
        return ws

    # Resumo
    ws = wb.active
    ws.title = S["summary"]
    ws.append([S["title"].format(pain=scope_of(meta))])
    ws["A1"].font = Font(bold=True, size=14)
    if meta.get("demo"):
        ws.append([S["demo"]])
    if meta.get("mode") == "hypothesis":
        ws.append([S["hypothesis"]])
    ws.append([S["date"], meta["generated_at"]])
    ws.append([S["market"], ", ".join(meta["market"]) if isinstance(meta.get("market"), list) else str(meta.get("market", ""))])
    ws.append([S["mode"], S["modes"][meta["mode"]]])
    ws.append([])
    for i, p in enumerate(products, 1):
        score = f" — {S['score']}: {p['traction_score']}/100" if "traction_score" in p else ""
        ws.append([f"{i}. {p['name']}", f"{p['solution_type']}{score}"])
    if report.get("verdict"):
        ws.append([])
        ws.append([S["verdict"], report["verdict"]])
    ws.append([])
    for a in list(report.get("assumptions") or []) + list(report.get("limits") or []):
        ws.append([S["assumptions"], a])
    ws.append([])
    ws.append([S["legend"]])
    ws.append([S["disclaimer"]])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 90
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    # Ranking
    sheet(
        "Ranking",
        ["#", S["fields"]["name"], S["fields"]["solution_type"], S["score"], S["fields"]["engagement"],
         S["fields"]["search_channel"], S["fields"]["traction_country"], S["fields"]["supplier"]],
        [
            [i, p["name"], p["solution_type"], p.get("traction_score", ""), fmt_value(p["engagement"]["value"]),
             fmt_value(p["search_channel"]["value"]), fmt_value(p["traction_country"]["value"]), fmt_value(p["supplier"]["value"])]
            for i, p in enumerate(products, 1)
        ],
        [5, 28, 24, 14, 36, 30, 22, 30],
    )

    # Fichas: uma linha por produto, um campo por coluna
    heads = [S["fields"]["name"], S["fields"]["pain"], S["fields"]["solution_type"]] + [S["fields"][k] for k in FIELD_KEYS]
    rows = [[p["name"], p["pain"], p["solution_type"]] + [fmt_value(p[k]["value"]) for k in FIELD_KEYS] for p in products]
    sheet("Fichas" if meta["language"] == "pt-BR" else "Profiles", heads, rows, [26, 24, 22] + [30] * len(FIELD_KEYS))

    # Evidências: formato longo, filtrável
    ev = evidence_rows(report)
    sheet(
        "Evidências" if meta["language"] == "pt-BR" else "Evidence",
        [S["fields"]["name"], "Campo" if meta["language"] == "pt-BR" else "Field", "Valor" if meta["language"] == "pt-BR" else "Value",
         S["src_cols"][2], S["src_cols"][0], S["src_cols"][1], "Nota" if meta["language"] == "pt-BR" else "Note",
         "Método" if meta["language"] == "pt-BR" else "Method"],
        [[r["product"], r["field"], r["value"], r["tag"], r["source"], r["date"], r["note"], r["method"]] for r in ev],
        [26, 32, 50, 16, 40, 12, 40, 18],
    )

    # Unit economics com FÓRMULAS (o usuário pode editar preço/custo/taxas)
    with_econ = [p for p in products if isinstance(p.get("economics"), dict)]
    if with_econ:
        E = S["econ"]
        heads = [S["fields"]["name"], E["price"], "Custo" if meta["language"] == "pt-BR" else "Cost",
                 "Frete" if meta["language"] == "pt-BR" else "Shipping", "Impostos" if meta["language"] == "pt-BR" else "Taxes",
                 "Taxa %" if meta["language"] == "pt-BR" else "Fee %", "Taxa fixa" if meta["language"] == "pt-BR" else "Fixed fee",
                 "Reembolso %" if meta["language"] == "pt-BR" else "Refund %", E["cost"], E["fees"], E["refund"],
                 E["profit"], E["margin"], E["cpa"], E["roas"]]
        ws = sheet(S["economics"], heads, [], [26] + [14] * 14)
        for r, p in enumerate(with_econ, start=2):
            e = p["economics"]
            ws.append([
                p["name"], e["price"], e["cost"], e.get("shipping", 0), e.get("tax", 0), e.get("fee_pct", 0),
                e.get("fee_fixed", 0), e.get("refund_pct", 0),
                f"=C{r}+D{r}+E{r}", f"=B{r}*F{r}/100+G{r}", f"=B{r}*H{r}/100", f"=B{r}-I{r}-J{r}-K{r}",
                f"=IF(B{r}>0,L{r}/B{r},0)", f"=MAX(L{r},0)", f'=IF(L{r}>0,B{r}/L{r},"{E["none"]}")',
            ])
            ws[f"M{r}"].number_format = "0.0%"
            for col in "BCDEGIJKLN":
                ws[f"{col}{r}"].number_format = "#,##0.00"
            ws[f"O{r}"].number_format = "0.00"

    # Fontes
    srcs = collect_sources(report)
    if srcs:
        sheet(
            S["sources"],
            list(S["src_cols"]),
            [[s["source"], s["date"], S["tags"].get(s["label"], s["label"]), "; ".join(s["used_in"])] for s in srcs],
            [60, 12, 16, 60],
        )
    wb.save(path)


# --------------------------------------------------------------------------- PDF

_PDF_MAP = str.maketrans({"−": "-", "⚠": "!", "✅": "", "🟡": "", "⚪": "", "’": "'", "→": "->"})


def _pdf_text(s: Any) -> str:
    """O PDF usa fontes padrão (WinAnsi): troca o que ficaria fora do alfabeto latino."""
    s = str(s).translate(_PDF_MAP)
    return s.encode("cp1252", errors="replace").decode("cp1252")


def to_pdf(blocks: list[dict[str, Any]], path: str) -> None:
    _need("reportlab", "reportlab", "pdf")
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#57606a"))
    cell = ParagraphStyle("cell", parent=body, fontSize=7.5, leading=9.5)
    cellh = ParagraphStyle("cellh", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    banner = ParagraphStyle("banner", parent=body, textColor=colors.HexColor("#b03000"), fontName="Helvetica-Bold")
    page = landscape(A4)
    margin = 1.5 * cm
    avail = page[0] - 2 * margin
    esc = lambda x: escape(_pdf_text(x))  # noqa: E731

    story: list[Any] = []
    for b in blocks:
        t = b["t"]
        if t == "title":
            story += [Paragraph(esc(b["text"]), ss["Title"]), Spacer(1, 6)]
        elif t == "banner":
            story += [Paragraph("! " + esc(b["text"]), banner), Spacer(1, 6)]
        elif t == "meta":
            story += [Paragraph("  |  ".join(f"<b>{esc(k)}:</b> {esc(v)}" for k, v in b["items"]), body), Spacer(1, 6)]
        elif t == "h2":
            story += [Spacer(1, 8), Paragraph(esc(b["text"]), ss["Heading2"])]
        elif t == "h3":
            story += [Paragraph(esc(b["text"]), ss["Heading3"])]
        elif t == "p":
            story += [Paragraph(esc(b["text"]), body), Spacer(1, 3)]
        elif t == "bullets":
            story += [Paragraph("&bull; " + esc(x), body) for x in b["items"]] + [Spacer(1, 4)]
        elif t == "ol":
            story += [Paragraph(f"{i}. " + esc(x), body) for i, x in enumerate(b["items"], 1)] + [Spacer(1, 4)]
        elif t == "note":
            story += [Paragraph(f"<i>{esc(b['text'])}</i>", small), Spacer(1, 4)]
        elif t == "field":
            color = "#%02x%02x%02x" % TAG_RGB[b["tag"]] if b.get("tag") else None
            tag = f" <font size=7.5 color='{color}'><b>[{esc(b['tagtext'])}]</b></font>" if color else ""
            src = f" <font size=7.5 color='#6e7781'>({esc(b['src'])})</font>" if b.get("src") else ""
            story.append(Paragraph(f"<b>{esc(b['label'])}:</b> {esc(b['value'])}{tag}{src}", body))
        elif t == "table":
            n = len(b["headers"])
            data = [[Paragraph(esc(h), cellh) for h in b["headers"]]]
            data += [[Paragraph(esc(c), cell) for c in r] for r in b["rows"]]
            tbl = Table(data, colWidths=[avail / n] * n, repeatRows=1)
            tbl.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d0d7de")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f8fa")]),
                    ]
                )
            )
            story += [Spacer(1, 4), tbl, Spacer(1, 8)]
    doc = SimpleDocTemplate(path, pagesize=page, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin)
    doc.build(story)
