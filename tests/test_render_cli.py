import contextlib
import io
import json
import os
import re
import tempfile
import unittest

from _helpers import EXAMPLE_BRIEF, EXAMPLE_REPORT, ROOT, example_report, has
import teixugo
from tx import blocks, office, render


def run(*argv):
    """Executa o CLI e devolve (código, saída)."""
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        try:
            code = teixugo.main(list(argv))
        except SystemExit as exc:  # erros de uso e sys.exit("mensagem")
            code = exc.code if isinstance(exc.code, int) else 1
            if isinstance(exc.code, str):
                buf.write(exc.code)
    return code, buf.getvalue()


class TextRenderers(unittest.TestCase):
    def setUp(self):
        self.report = example_report()
        self.blocks = blocks.build_blocks(self.report)

    def test_markdown_has_all_requested_fields(self):
        md = render.to_md(self.blocks)
        for label in (
            "Produto", "Dor", "Segmento", "Público-alvo", "Efetividade", "Engajamento", "Link (maior engajamento)",
            "Redes sociais com maior engajamento", "Canal com maior busca", "Fornecedor", "País de maior tração",
            "País do fornecedor", "Primeira aparição verificada", "Comparação",
        ):
            self.assertIn(f"**{label}:**", md)
        self.assertIn("✅", md)
        self.assertIn("EXEMPLO FICTÍCIO", md)
        self.assertIn("## Comparação", md)
        self.assertIn("## Unit economics", md)

    def test_markdown_numbering_is_not_duplicated(self):
        md = render.to_md(self.blocks)
        self.assertNotIn("- 1.", md)

    def test_txt_has_no_markdown_syntax(self):
        txt = render.to_txt(self.blocks)
        self.assertNotIn("**", txt)
        self.assertIn("[VERIFICADO]", txt)

    def test_html_is_self_contained_and_escaped(self):
        self.report["products"][0]["name"] = "Luva <script>alert(1)</script>"
        html = render.to_html(blocks.build_blocks(self.report))
        self.assertIn("<!doctype html>", html)
        self.assertNotIn("<script>", html)
        self.assertNotIn("http://", html.replace("https://", ""))

    def test_english_report_uses_english_labels(self):
        self.report["meta"]["language"] = "en"
        md = render.to_md(blocks.build_blocks(self.report))
        self.assertIn("**Supplier:**", md)
        self.assertIn("**Top search channel:**", md)
        self.assertIn("## Comparison", md)

    def test_hypothesis_banner(self):
        self.report["meta"]["mode"] = "hypothesis"
        md = render.to_md(blocks.build_blocks(self.report))
        self.assertIn("MODO HIPÓTESE", md)

    def test_pipe_in_value_does_not_break_table(self):
        self.report["products"][0]["supplier"]["value"] = "A | B"
        md = render.to_md(blocks.build_blocks(self.report))
        self.assertIn("A \\| B", md)


@unittest.skipUnless(has("docx") and has("openpyxl") and has("reportlab"), "bibliotecas opcionais ausentes")
class OfficeRenderers(unittest.TestCase):
    def setUp(self):
        self.report = example_report()
        self.blocks = blocks.build_blocks(self.report)
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def test_docx_opens_and_has_content(self):
        import docx

        path = os.path.join(self.tmp.name, "r.docx")
        office.to_docx(self.blocks, path)
        text = "\n".join(p.text for p in docx.Document(path).paragraphs)
        self.assertIn("Luva removedora", text)
        self.assertIn("Efetividade", text)

    def test_xlsx_has_sheets_and_formulas(self):
        import openpyxl

        path = os.path.join(self.tmp.name, "r.xlsx")
        office.to_xlsx(self.report, path)
        wb = openpyxl.load_workbook(path)
        for name in ("Resumo", "Ranking", "Fichas", "Evidências", "Unit economics", "Fontes"):
            self.assertIn(name, wb.sheetnames)
        ws = wb["Unit economics"]
        self.assertEqual(ws["L2"].value, "=B2-I2-J2-K2")  # lucro antes de anúncio como fórmula
        self.assertTrue(str(ws["O2"].value).startswith("=IF("))

    def test_pdf_is_valid(self):
        path = os.path.join(self.tmp.name, "r.pdf")
        office.to_pdf(self.blocks, path)
        with open(path, "rb") as fh:
            self.assertEqual(fh.read(5), b"%PDF-")


class CLI(unittest.TestCase):
    def test_validate_example(self):
        code, out = run("validate", EXAMPLE_REPORT)
        self.assertEqual(code, 0)
        self.assertIn("válido", out)

    def test_validate_reports_errors_with_exit_2(self):
        r = example_report()
        del r["products"][0]["engagement"]["source"]
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "bad.json")
            with open(path, "w", encoding="utf-8") as fh:
                json.dump(r, fh)
            code, out = run("validate", path)
        self.assertEqual(code, 2)
        self.assertIn("ERRO", out)

    def test_render_blocks_invalid_report(self):
        r = example_report()
        r["meta"]["mode"] = "hypothesis"
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "bad.json")
            with open(path, "w", encoding="utf-8") as fh:
                json.dump(r, fh)
            code, _ = run("render", path, "--format", "md", "--out", os.path.join(tmp, "out"))
            self.assertEqual(code, 2)
            self.assertFalse(os.path.exists(os.path.join(tmp, "out")))

    def test_render_md_txt_html(self):
        with tempfile.TemporaryDirectory() as tmp:
            code, out = run("render", EXAMPLE_REPORT, "--format", "md,txt,html", "--out", tmp)
            self.assertEqual(code, 0, out)
            files = sorted(os.listdir(tmp))
            self.assertEqual(files, [
                "remover-pelo-de-cachorro-2026-10-05.html",
                "remover-pelo-de-cachorro-2026-10-05.md",
                "remover-pelo-de-cachorro-2026-10-05.txt",
            ])

    def test_render_unknown_format(self):
        code, out = run("render", EXAMPLE_REPORT, "--format", "gif")
        self.assertEqual(code, 2)
        self.assertIn("desconhecido", out)

    def test_rank_write_roundtrip(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "r.json")
            with open(EXAMPLE_REPORT, encoding="utf-8") as src, open(path, "w", encoding="utf-8") as dst:
                dst.write(src.read())
            code, _ = run("rank", path, "--write")
            self.assertEqual(code, 0)
            with open(path, encoding="utf-8") as fh:
                data = json.load(fh)
            self.assertIn("traction_score", data["products"][0])
            self.assertEqual(run("validate", path)[0], 0)

    def test_queries_and_brief_commands(self):
        code, out = run("queries", EXAMPLE_BRIEF, "--format", "json", "--limit", "5")
        self.assertEqual(code, 0)
        self.assertEqual(len(json.loads(out)), 5)
        self.assertEqual(run("validate-brief", EXAMPLE_BRIEF)[0], 0)
        code, out = run("brief-template")
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(out)["pain"], "remover pelo de cachorro")

    def test_economics_command(self):
        code, out = run("economics", "--price", "100", "--cost", "30", "--shipping", "10")
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(out)["base"]["profit_before_ads"], 60.0)

    def test_check_links_skips_demo(self):
        code, out = run("check-links", EXAMPLE_REPORT)
        self.assertEqual(code, 0)
        self.assertIn("demonstração", out)

    def test_missing_file(self):
        code, out = run("validate", os.path.join(ROOT, "nao-existe.json"))
        self.assertNotEqual(code, 0)
        self.assertIn("não encontrado", out)


class SkillFile(unittest.TestCase):
    """O SKILL.md precisa ser carregável e não pode apontar para arquivos que não existem."""

    def setUp(self):
        with open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8") as fh:
            self.text = fh.read()

    def test_frontmatter(self):
        m = re.match(r"^---\n(.*?)\n---\n", self.text, re.S)
        self.assertIsNotNone(m, "SKILL.md precisa começar com frontmatter YAML")
        fm = m.group(1)
        name = re.search(r"^name:\s*(.+)$", fm, re.M).group(1).strip()
        desc = re.search(r"^description:\s*(.+)$", fm, re.M).group(1).strip()
        self.assertEqual(name, "teixugo-minerador")
        self.assertRegex(name, r"^[a-z0-9]+(-[a-z0-9]+)*$")
        self.assertLessEqual(len(name), 64)
        self.assertLessEqual(len(desc), 1024)
        self.assertGreater(len(desc), 100)
        # YAML: um "dois-pontos + espaço" ou " #" num valor sem aspas quebra o frontmatter
        self.assertNotIn(": ", desc)
        self.assertNotIn(" #", desc)

    def test_size_budget(self):
        self.assertLess(len(self.text.splitlines()), 500)

    def test_referenced_files_exist(self):
        refs = set(re.findall(r"(?:references|scripts|examples)/[\w./-]+\.(?:md|py|json)", self.text))
        self.assertTrue(refs)
        for ref in refs:
            self.assertTrue(os.path.exists(os.path.join(ROOT, ref)), f"referência quebrada no SKILL.md: {ref}")


if __name__ == "__main__":
    unittest.main()
